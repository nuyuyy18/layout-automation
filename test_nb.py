import docx
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_test():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(14.8), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(1.8)

    # Add notebook section
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    c.width = Cm(10.8)

    tcPr = c._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FDFDFB"/>'))
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="140" w:type="dxa"/><w:bottom w:w="140" w:type="dxa"/>'
        f'<w:left w:w="160" w:type="dxa"/><w:right w:w="160" w:type="dxa"/></w:tcMar>'
    ))
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    borders_xml += '<w:top w:val="single" w:sz="12" w:space="0" w:color="1B5E20"/>'
    borders_xml += '<w:left w:val="single" w:sz="18" w:space="0" w:color="1B5E20"/>'
    borders_xml += '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="1B5E20"/>'
    borders_xml += '<w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

    p0 = c.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run('✍️  MY SACRED SOUL REFLECTION NOTEBOOK')
    r0.font.name = 'Garamond'
    r0.font.size = Pt(11.5)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(27, 94, 32)

    pq = c.add_paragraph()
    pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pq.paragraph_format.space_after = Pt(8)
    rq = pq.add_run('"Tuliskan apa yang membuat jiwamu bergetar dalam hening di hadapan Allah."')
    rq.font.name = 'Garamond'
    rq.font.size = Pt(9.5)
    rq.font.italic = True
    rq.font.color.rgb = RGBColor(90, 90, 90)

    # Date header with dot leader edge-to-edge
    p_date = c.add_paragraph()
    p_date.paragraph_format.tab_stops.add_tab_stop(Cm(10.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p_date.paragraph_format.space_after = Pt(4)
    rd = p_date.add_run('Tanggal / Pekan ke : ')
    rd.font.name = 'Garamond'
    rd.font.size = Pt(9.5)
    rd.font.bold = True
    p_date.add_run('\t')

    p_mood = c.add_paragraph()
    p_mood.paragraph_format.space_after = Pt(10)
    rm = p_mood.add_run('Kondisi Batin Hari Ini :   [  ] Tenang     [  ] Lelah     [  ] Gundah     [  ] Bersyukur')
    rm.font.name = 'Garamond'
    rm.font.size = Pt(9)
    rm.font.italic = True
    rm.font.color.rgb = RGBColor(70, 70, 70)

    prompts = [
        "1. Hal yang paling membuat batinku merasa sepi / cemas pekan ini:",
        "2. Situasi atau tuntutan yang paling menguras energiku (tanda burnout):",
        "3. Batasan diri (boundary) yang berhasil kutegakkan demi kesehatan batin:",
        "4. Satu hal sederhana yang paling kusyukuri dari nikmat Allah pekan ini:"
    ]

    for q in prompts:
        p_q = c.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        rq = p_q.add_run(q)
        rq.font.name = 'Garamond'
        rq.font.size = Pt(9.5)
        rq.font.bold = True
        rq.font.color.rgb = RGBColor(27, 94, 32)

        for _ in range(2):
            pl = c.add_paragraph()
            pl.paragraph_format.tab_stops.add_tab_stop(Cm(10.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            pl.paragraph_format.space_before = Pt(2)
            pl.paragraph_format.space_after = Pt(2)
            pl.paragraph_format.line_spacing = 1.3
            rl = pl.add_run('\t')
            rl.font.name = 'Garamond'
            rl.font.size = Pt(10)
            rl.font.color.rgb = RGBColor(160, 160, 160)

    p_foot = c.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(10)
    p_foot.paragraph_format.space_after = Pt(2)
    rf = p_foot.add_run('Afirmasi Jiwa: "Aku berhak tenang dan aku memilih berserah sepenuhnya kepada Allah."')
    rf.font.name = 'Garamond'
    rf.font.size = Pt(9)
    rf.font.italic = True
    rf.font.color.rgb = RGBColor(80, 80, 80)

    p_sign = c.add_paragraph()
    p_sign.paragraph_format.tab_stops.add_tab_stop(Cm(10.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p_sign.paragraph_format.space_before = Pt(4)
    rs = p_sign.add_run('Tanda Tangan Jiwa : ')
    rs.font.name = 'Garamond'
    rs.font.size = Pt(9.5)
    rs.font.bold = True
    p_sign.add_run('\t')

    doc.save('test_nb_result.docx')
    print('SUCCESS')

if __name__ == '__main__':
    create_test()
