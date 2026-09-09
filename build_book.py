import os
import json
import re
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Palet Warna Novel Digital — elegan, monokrom, bersih
C_BLACK  = RGBColor(20,  20,  20)
C_DARK   = RGBColor(45,  45,  45)
C_MUTED  = RGBColor(90,  90,  90)
C_LIGHT  = RGBColor(145, 145, 145)
C_ACCENT = RGBColor(55,  55,  55)

FONT_BODY = "Garamond"

def clean_all(text):
    if not text:
        return ""
    t = text.replace("\u2014", ", ").replace("\u2013", " ")
    t = re.sub(r'tere\s*liye,?\s*', '', t, flags=re.IGNORECASE)
    t = t.replace(", ,", ",").replace(" ,", ",")
    while "  " in t:
        t = t.replace("  ", " ")
    return t.strip()

def add_p(doc, text="", bold=False, italic=False, size=11.5,
          space_after=6, space_before=0, line_spacing=1.35,
          indent=0.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=C_BLACK):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent > 0:
        p.paragraph_format.first_line_indent = Cm(indent)
    if text:
        cleaned = clean_all(text)
        r = p.add_run(cleaned)
        r.font.name   = FONT_BODY
        r.font.size   = Pt(size)
        r.font.bold   = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return p

def add_thin_rule(doc, space_before=8, space_after=8):
    """Garis pemisah tipis — bersih, tanpa simbol dekoratif."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\u2500" * 28)
    r.font.name      = FONT_BODY
    r.font.size      = Pt(8)
    r.font.color.rgb = C_LIGHT

def add_block_quote(doc, text, size=10.5):
    """Kutipan / refleksi — indented sederhana, tanpa border tabel."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(10)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.left_indent  = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.8)
    r = p.add_run(clean_all(text))
    r.font.name   = FONT_BODY
    r.font.size   = Pt(size)
    r.font.italic = True
    r.font.color.rgb = C_MUTED

def add_arabic_block(doc, ayat_ar, ayat_id, ayat_ref,
                     doa_ar=None, doa_id=None, doa_ref=None):
    """Ayat & doa sebagai blok teks biasa — tanpa tabel / kotak."""
    for ar, trans, ref in [(ayat_ar, ayat_id, ayat_ref),
                           (doa_ar,  doa_id,  doa_ref)]:
        if ar:
            p_ar = doc.add_paragraph()
            p_ar.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_ar.paragraph_format.space_before = Pt(10)
            p_ar.paragraph_format.space_after  = Pt(4)
            p_ar.paragraph_format.line_spacing = 1.8
            r_ar = p_ar.add_run(ar.strip())
            r_ar.font.name      = "Traditional Arabic"
            r_ar.font.size      = Pt(15)
            r_ar.font.bold      = True
            r_ar.font.color.rgb = C_BLACK
        if trans:
            p_tr = doc.add_paragraph()
            p_tr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_tr.paragraph_format.space_before = Pt(2)
            p_tr.paragraph_format.space_after  = Pt(2)
            p_tr.paragraph_format.line_spacing = 1.25
            p_tr.paragraph_format.left_indent  = Cm(0.5)
            p_tr.paragraph_format.right_indent = Cm(0.5)
            r_tr = p_tr.add_run(f'"{clean_all(trans)}"')
            r_tr.font.name      = FONT_BODY
            r_tr.font.size      = Pt(10.5)
            r_tr.font.italic    = True
            r_tr.font.color.rgb = C_MUTED
        if ref:
            p_ref = doc.add_paragraph()
            p_ref.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_ref.paragraph_format.space_before = Pt(0)
            p_ref.paragraph_format.space_after  = Pt(8)
            r_ref = p_ref.add_run(f"({clean_all(ref)})")
            r_ref.font.name      = FONT_BODY
            r_ref.font.size      = Pt(9)
            r_ref.font.bold      = True
            r_ref.font.color.rgb = C_ACCENT

def generate_book(json_path="book_content.json",
                  out_docx="result/Ketika_Jiwa_Pulang_ke_Sunyi.docx"):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width,   sec.page_height   = Cm(14.8), Cm(21.0)
    sec.top_margin,   sec.bottom_margin = Cm(2.2),  Cm(2.2)
    sec.left_margin,  sec.right_margin  = Cm(2.4),  Cm(2.0)
    sec.different_first_page_header_footer = True

    # Nomor halaman — hanya angka di tengah footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>'))

    fm = data["front_matter"]

    # ── 1. HALAMAN JUDUL ──────────────────────────────────────────────────────
    doc.add_paragraph().paragraph_format.space_before = Pt(60)
    add_p(doc, fm["title"], bold=True, size=22,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=10)
    add_p(doc, fm["subtitle"], italic=True, size=11,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0, space_after=30)
    add_p(doc, fm["compiler"], size=10,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0, space_after=4)
    for insp in fm["inspirations"]:
        add_p(doc, insp, italic=True, size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=C_LIGHT, indent=0, space_after=2)
    doc.add_paragraph().paragraph_format.space_before = Pt(50)
    add_p(doc, fm["edition"], size=9,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_LIGHT, indent=0, space_after=2)
    add_p(doc, f"Yogyakarta, {fm['year']}", size=9,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_LIGHT, indent=0)

    # ── 2. KOLOFON ────────────────────────────────────────────────────────────
    doc.add_page_break()
    doc.add_paragraph().paragraph_format.space_before = Pt(40)
    add_p(doc, fm["title"], bold=True, size=15,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=4)
    add_p(doc, fm["subtitle"], italic=True, size=10,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0, space_after=28)
    add_p(doc, "Hak Cipta & Keterangan Penerbitan", bold=True, size=9.5,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_ACCENT, indent=0, space_after=8)
    add_p(doc, "Buku ini disusun sebagai materi refleksi spiritual dan psikologi islami "
               "berbasis riset ilmiah internasional.",
          italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
          color=C_MUTED, indent=0, space_after=4)
    add_p(doc, "Format Buku: Ukuran Standar A5 (14.8 cm \u00d7 21.0 cm)",
          size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0, space_after=4)
    add_p(doc, f"Cetakan Pertama: September {fm['year']}",
          size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0)

    # ── 3. EPIGRAF ────────────────────────────────────────────────────────────
    doc.add_page_break()
    doc.add_paragraph().paragraph_format.space_before = Pt(80)
    author_text = clean_all(fm["epigraph"].get("author", "Catatan Perjalanan Batin"))
    add_block_quote(doc,
        fm["epigraph"]["quote"] + f"\n\n\u2014 {author_text}",
        size=11)

    # ── 4. PROLOG ─────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_p(doc, fm["prologue"]["title"], bold=True, size=15,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=18)
    for p in fm["prologue"]["paragraphs"]:
        add_p(doc, p)

    # ── 5. DAFTAR ISI ─────────────────────────────────────────────────────────
    doc.add_page_break()
    add_p(doc, "DAFTAR ISI", bold=True, size=14,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_ACCENT, indent=0, space_after=4)
    add_p(doc, "Peta Perjalanan Jiwa Menuju Ketenangan Sejati",
          italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER,
          color=C_MUTED, indent=0, space_after=18)
    for part in data["parts"]:
        add_p(doc, f"{part['num']}  \u2014  {part['title']}", bold=True, size=10.5,
              color=C_ACCENT, indent=0, space_after=2)
        for chap in part["chapters"]:
            add_p(doc, f"    {chap['num']}:  {chap['title']}", size=9.5,
                  indent=0.6, space_after=2, color=C_DARK)
    add_p(doc, "Epilog: Sebuah Surat untuk Jiwa yang Berjuang",
          bold=True, size=10, indent=0, space_after=2)

    # ── 6. BAGIAN & BAB ───────────────────────────────────────────────────────
    for part in data["parts"]:
        doc.add_page_break()
        doc.add_paragraph().paragraph_format.space_before = Pt(72)
        add_p(doc, part["num"].upper(), bold=True, size=11,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=C_LIGHT, indent=0, space_after=8)
        add_p(doc, part["title"], bold=True, size=18,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=12)
        add_p(doc, part["subtitle"], italic=True, size=10.5,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0, space_after=16)
        add_p(doc, f"Inspirasi Kajian: {part['source']}", italic=True, size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=C_LIGHT, indent=0)

        for chap in part["chapters"]:
            doc.add_page_break()

            # Judul bab — bersih, tanpa ❖
            add_p(doc, chap["num"].upper(), bold=True, size=10.5,
                  align=WD_ALIGN_PARAGRAPH.CENTER, color=C_LIGHT, indent=0, space_after=4)
            add_p(doc, chap["title"], bold=True, size=16,
                  align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=6)
            if chap.get("subtitle"):
                add_p(doc, chap["subtitle"], italic=True, size=10.5,
                      align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0, space_after=16)

            # Ayat & Doa — blok teks biasa, tanpa tabel kotak hijau
            add_arabic_block(
                doc,
                chap.get("ayat_ar"), chap.get("ayat_id"), chap.get("ayat_ref"),
                chap.get("doa_ar"),  chap.get("doa_id"),  chap.get("doa_ref"),
            )

            # Narasi utama (catatan hati) — langsung teks, tanpa header label [ ... ]
            for p in chap["catatan_hati"]:
                add_p(doc, p)

            # Paragraf sains — langsung teks, tanpa label header [ Sains ... ]
            doc.add_paragraph().paragraph_format.space_before = Pt(6)
            for p in chap["sains_paras"]:
                add_p(doc, p, size=11, space_after=5)

            # Kutipan hikmah — block quote tipis
            add_block_quote(doc, f'"{chap["quote"]}"')

    # ── 7. EPILOG ─────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_p(doc, data["epilogue"]["title"], bold=True, size=15,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=18)
    for p in data["epilogue"]["paragraphs"]:
        add_p(doc, p)

    # ── 8. RISET ILMIAH ───────────────────────────────────────────────────────
    doc.add_page_break()
    add_p(doc, "POJOK RISET & VALIDASI ILMIAH", bold=True, size=12,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=C_ACCENT, indent=0, space_after=4)
    add_p(doc, "Daftar Sitasi Jurnal Peer-Reviewed dengan DOI Aktif",
          italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER,
          color=C_MUTED, indent=0, space_after=16)
    for idx, r in enumerate(data["research_data"], start=1):
        add_p(doc, f"[{idx}] {clean_all(r['citation'])}", bold=True, size=9.5,
              space_after=1, indent=0)
        add_p(doc, f"DOI: {r['doi']}", size=8.5, color=C_MUTED, space_after=1, indent=0)
        add_p(doc, clean_all(r["notes"]), italic=True, size=9,
              color=C_MUTED, space_after=8, indent=0)

    # ── 10. QR CODE ───────────────────────────────────────────────────────────
    res_dir = os.path.dirname(out_docx)
    for i, qr in enumerate(data["digital_integration"]["qr_items"]):
        doc.add_page_break()
        if i == 0:
            add_p(doc, "INTEGRASI DIGITAL & KODE QR PENDUKUNG", bold=True, size=12,
                  align=WD_ALIGN_PARAGRAPH.CENTER, color=C_ACCENT, indent=0, space_after=4)
        add_p(doc, clean_all(qr["title"]), bold=True, size=11,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=4)
        add_p(doc, clean_all(qr["desc"]), italic=True, size=10,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=C_MUTED, indent=0, space_after=4)
        add_p(doc, f"Tautan: {qr['url']}", size=9,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=C_LIGHT, indent=0, space_after=10)
        img_path = os.path.join(res_dir, qr["img_file"])
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after  = Pt(8)
            p_img.add_run().add_picture(img_path, width=Inches(1.4))

    os.makedirs(os.path.dirname(out_docx), exist_ok=True)
    doc.save(out_docx)
    print("SUCCESS: Saved", out_docx)


if __name__ == "__main__":
    generate_book()
